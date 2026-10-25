import os
import sys
import uuid
import argparse
import zipfile
import shutil
import json

try:
    import pandas as pd
except ImportError:
    pd = None
    print("[WARN] Pandas not installed. CSV/XLS extraction will fail. Run: pip install pandas openpyxl")

try:
    import pdfplumber
except ImportError:
    pdfplumber = None
    print("[WARN] pdfplumber not installed. PDF extraction will fail. Run: pip install pdfplumber")

try:
    from docx import Document as DocxDocument
except ImportError:
    DocxDocument = None
    print("[WARN] python-docx not installed. DOCX extraction will fail. Run: pip install python-docx")

try:
    from pptx import Presentation as PptxPresentation
except ImportError:
    PptxPresentation = None
    print("[WARN] python-pptx not installed. PPTX extraction will fail. Run: pip install python-pptx")

try:
    from PIL import Image
    import pytesseract
    _ocr_available = True
except ImportError:
    _ocr_available = False
    print("[WARN] Pillow/pytesseract not installed. Image OCR will be skipped. Run: pip install Pillow pytesseract")

# Define core paths
BASE_DIR = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
RAW_DIR = os.path.join(BASE_DIR, ".raw")
ARCHIVE_DIR = os.path.join(RAW_DIR, "archive")
WIKI_DIR = os.path.join(BASE_DIR, "wiki")
ENTITIES_DIR = os.path.join(WIKI_DIR, "entities")
META_DIR = os.path.join(WIKI_DIR, "meta")
INDEX_PATH = os.path.join(META_DIR, "index.md")

os.makedirs(ENTITIES_DIR, exist_ok=True)
os.makedirs(META_DIR, exist_ok=True)
os.makedirs(ARCHIVE_DIR, exist_ok=True)

def update_index(entity_name, filepath, format_type=".md"):
    """Updates the central index.md catalog with the new node."""
    entry = f"- [{entity_name}](../entities/{os.path.basename(filepath)})\n"
    mode = 'a' if os.path.exists(INDEX_PATH) else 'w'
    with open(INDEX_PATH, mode) as f:
        if mode == 'w':
            f.write("# Knowledge Base Index\n\n## Entities & Concepts\n")
        f.write(entry)
    print(f"[INDEX] Updated index.md with {entity_name}")

def extract_pdf_text(filepath):
    if not pdfplumber:
        raise ImportError("pdfplumber is required for PDFs.")
    full_text = []
    with pdfplumber.open(filepath) as pdf:
        for page in pdf.pages:
            text = page.extract_text()
            if text:
                full_text.append(text)
    return "\n".join(full_text)

def extract_docx_text(filepath):
    if not DocxDocument:
        raise ImportError("python-docx is required for DOCX files.")
    doc = DocxDocument(filepath)
    paragraphs = [p.text for p in doc.paragraphs if p.text.strip()]
    # Also extract text from tables
    for table in doc.tables:
        for row in table.rows:
            row_text = " | ".join(cell.text.strip() for cell in row.cells if cell.text.strip())
            if row_text:
                paragraphs.append(row_text)
    return "\n".join(paragraphs)

def extract_pptx_text(filepath):
    if not PptxPresentation:
        raise ImportError("python-pptx is required for PPTX files.")
    prs = PptxPresentation(filepath)
    slides_text = []
    for i, slide in enumerate(prs.slides, 1):
        slide_parts = [f"[Slide {i}]"]
        for shape in slide.shapes:
            if hasattr(shape, "text") and shape.text.strip():
                slide_parts.append(shape.text.strip())
        if len(slide_parts) > 1:
            slides_text.append("\n".join(slide_parts))
    return "\n\n".join(slides_text)

def extract_image_text(filepath):
    if not _ocr_available:
        return f"[Image file: {os.path.basename(filepath)} — OCR unavailable, install pytesseract + tesseract-ocr]"
    try:
        img = Image.open(filepath)
        # Try Polish+English first (corpus is PL), fall back to eng-only
        try:
            text = pytesseract.image_to_string(img, lang="pol+eng")
        except pytesseract.TesseractError:
            text = pytesseract.image_to_string(img, lang="eng")
        return text.strip() if text.strip() else f"[Image: {os.path.basename(filepath)} — no readable text detected by OCR]"
    except Exception as e:
        return f"[Image extraction failed: {e}]"

def extract_audio_text(filepath):
    """Hits ElevenLabs API to dump an audio file into a raw text transcript."""
    import requests
    api_key = os.getenv("ELEVENLABS_API_KEY")
    if not api_key:
        print("[WARN] No ELEVENLABS_API_KEY set. Bypassing cloud transcription and returning mocked text.")
        return "Audio Transcript Placeholder: The speakers discussed the integration of legacy memory modules into autonomous workflows."
        
    print("[CLOUD] Sending audio payload to ElevenLabs transcription...")
    url = "https://api.elevenlabs.io/v1/speech-to-text"
    headers = {
        "xi-api-key": api_key,
        "Accept": "application/json"
    }
    try:
        with open(filepath, "rb") as audio_file:
            files = {"file": (os.path.basename(filepath), audio_file, "audio/mpeg")}
            response = requests.post(url, headers=headers, files=files)
        if response.status_code == 200:
            return response.json().get("text", "")
        else:
            return f"Transcript extraction failed due to API rejection: {response.text}"
    except Exception as e:
        return f"Transcript extraction failed: {e}"

def extract_google_doc(filepath):
    """Uses Google API to fetch the text content of a .gdoc file."""
    try:
        from google.oauth2 import service_account
        from googleapiclient.discovery import build
        
        with open(filepath, 'r') as f:
            gdoc_data = json.load(f)
            
        doc_id = gdoc_data.get("doc_id")
        if not doc_id:
            return "Invalid .gdoc format. No doc_id found."
            
        creds_path = os.path.join(BASE_DIR, "credentials.json")
        if not os.path.exists(creds_path):
            print("[WARN] credentials.json missing. Cannot authenticate to Google Workspace.")
            return "Google Auth required. Supply a credentials.json Service Account key in the root directory."
            
        creds = service_account.Credentials.from_service_account_file(
            creds_path, scopes=['https://www.googleapis.com/auth/documents.readonly']
        )
        service = build('docs', 'v1', credentials=creds)
        doc = service.documents().get(documentId=doc_id).execute()
        
        text_content = ""
        for element in doc.get('body').get('content', []):
            if 'paragraph' in element:
                for req in element.get('paragraph').get('elements', []):
                    if 'textRun' in req:
                        text_content += req.get('textRun').get('content')
        return text_content
    except ImportError:
        print("[WARN] Google API dependencies missing. Run: pip install -r requirements.txt")
        return "Google Workspace extraction failed due to missing pip packages."
    except Exception as e:
        return f"Google extraction failed: {e}"

def process_spreadsheet_to_xml(filepath):
    if not pd:
        raise ImportError("pandas is required for spreadsheets.")
    ext = os.path.splitext(filepath)[1].lower()
    if ext == ".csv":
        df = pd.read_csv(filepath)
    else:
        df = pd.read_excel(filepath)
    
    xml_output = [f'<?xml version="1.0" encoding="UTF-8"?>']
    base_name = os.path.basename(filepath)
    grid_id = f"grid_{uuid.uuid4().hex[:8]}"
    xml_output.append(f'<DataGrid id="{grid_id}" origin_source=".raw/{base_name}">')
    xml_output.append('    <Columns>')
    for col in df.columns:
        dtype = "number" if pd.api.types.is_numeric_dtype(df[col]) else "string"
        clean_name = str(col).replace("&", "and").replace("<", "").replace(">", "").strip()
        xml_output.append(f'        <Column type="{dtype}" name="{clean_name}" />')
    xml_output.append('    </Columns>')
    xml_output.append('    <Rows>')
    for _, row in df.iterrows():
        xml_output.append('        <Row>')
        for col in df.columns:
            clean_col = str(col).replace("&", "and").replace(" ", "_").replace("<", "").replace(">", "").strip()
            val = str(row[col]).replace("&", "&amp;").replace("<", "&lt;").replace(">", "&gt;")
            xml_output.append(f'            <{clean_col}>{val}</{clean_col}>')
        xml_output.append('        </Row>')
    xml_output.append('    </Rows>')
    xml_output.append('</DataGrid>')
    
    final_xml = "\n".join(xml_output)
    entity_name = base_name.split('.')[0].replace(" ", "_").capitalize()
    entity_filepath = os.path.join(ENTITIES_DIR, f"{entity_name}.xml")
    
    with open(entity_filepath, 'w', encoding='utf-8') as f:
        f.write(final_xml)
    print(f"[WIKI] Matrix routed to XML schema at {entity_filepath}")
    update_index(entity_name, entity_filepath, format_type=".xml")
    return True

def process_zip_archive(filepath, mock=False):
    """Extracts a zip file to a temporary folder and processes its contents, then archives the zip."""
    print(f"[INGEST] Unpacking ZIP Archive: {filepath}")
    extract_dir = filepath + "_unpacked"
    os.makedirs(extract_dir, exist_ok=True)
    
    try:
        with zipfile.ZipFile(filepath, 'r') as zip_ref:
            zip_ref.extractall(extract_dir)
            
        print("[INGEST] Recursively iterating through unpacked contents...")
        for root, _, files in os.walk(extract_dir):
            for file in files:
                if not file.startswith('.'):
                    unpacked_filepath = os.path.join(root, file)
                    ingest_file(unpacked_filepath, mock=mock)
                    
    except Exception as e:
        print(f"[ERR] Failed to complete ZIP extraction: {e}")
        
    print(f"[CLEANUP] Archiving original bundle {filepath}")
    shutil.move(filepath, os.path.join(ARCHIVE_DIR, os.path.basename(filepath)))

def mock_llm_parse(file_path, content):
    print(f"[LLM] MOCK Parsing abstract for {os.path.basename(file_path)}...")
    mock_entity_name = os.path.basename(file_path).split('.')[0].replace(" ", "_").capitalize()
    mock_yaml = f"---\nid: {str(uuid.uuid4())}\ntype: synthesis\ntags: [auto_generated, mock]\ninbound_links: 0\nconfidence: 0.90\nlast_ingested_sources: [\"{os.path.basename(file_path)}\"]\n---"
    sample_content = content[:150].replace('\n', ' ') if content else "No readable text extracted."
    mock_body = f"\n# Overview\nThis is a mocked output.\n\n> {sample_content}...\n\n## Citations\n- [../../.raw/{os.path.basename(file_path)}](../../.raw/{os.path.basename(file_path)})\n\n## Related Nodes\n- [[Placeholder]]"
    return mock_entity_name, mock_yaml, mock_body

def real_llm_parse(file_path, content):
    api_key = os.getenv("ANTHROPIC_API_KEY")
    if not api_key:
        base_name = os.path.basename(file_path)
        fallback_path = os.path.join(RAW_DIR, f"{base_name}_extracted.txt")
        print(f"[BASE PATH] No API key found. Dumping payload to {fallback_path} for CLI Agent to synthesize.")
        with open(fallback_path, 'w', encoding='utf-8') as f:
            f.write(content)
        return None, None, None
        
    from anthropic import Anthropic
    client = Anthropic(api_key=api_key)
    prompt = f'Extract standard slug from document printed as ENTITY_NAME: slug_name. Then Output Strict YAML and Markdown Body: "{content}"'

    try:
        response = client.messages.create(
            model="claude-opus-4-7",
            max_tokens=1500,
            messages=[{"role": "user", "content": prompt}]
        )
        return ("Extracted_Concept", "---\nid: generated\n---", response.content[0].text.strip())
    except Exception as e:
        print(f"[ERR] Anthropic API failed: {e}")
        return None, None, None

def ingest_file(file_path, mock=False):
    if not os.path.exists(file_path):
        print(f"Error: {file_path} doesn't exist.")
        return

    ext = os.path.splitext(file_path)[1].lower()
    
    if ext == ".zip":
        process_zip_archive(file_path, mock)
        return
        
    if ext in [".csv", ".xls", ".xlsx"]:
        print(f"[INGEST] Routing tabular data to XML parser: {file_path}")
        process_spreadsheet_to_xml(file_path)
        return

    raw_content = ""
    if ext == ".pdf":
        print(f"[INGEST] Extracting raw text from PDF: {file_path}")
        raw_content = extract_pdf_text(file_path)
    elif ext in [".mp3", ".wav", ".m4a"]:
        print(f"[INGEST] Found Audio file. Initiating Voice-to-Text: {file_path}")
        raw_content = extract_audio_text(file_path)
    elif ext == ".gdoc":
        print(f"[INGEST] Processing Google Workspace Document: {file_path}")
        raw_content = extract_google_doc(file_path)
    elif ext in [".txt", ".md"]:
        print(f"[INGEST] Reading standard text document: {file_path}")
        with open(file_path, 'r', encoding='utf-8') as f:
            raw_content = f.read()
    elif ext in [".docx", ".doc"]:
        print(f"[INGEST] Extracting text from Word document: {file_path}")
        raw_content = extract_docx_text(file_path)
    elif ext in [".pptx", ".ppt"]:
        print(f"[INGEST] Extracting slides from PowerPoint: {file_path}")
        raw_content = extract_pptx_text(file_path)
    elif ext in [".jpg", ".jpeg", ".png", ".gif", ".bmp", ".tiff", ".webp"]:
        print(f"[INGEST] Running OCR on image: {file_path}")
        raw_content = extract_image_text(file_path)
    else:
        print(f"[INGEST] Ignored unsupported file format: {ext}")
        return

    if mock:
        entity_name, yaml_frontmatter, body = mock_llm_parse(file_path, raw_content)
    else:
        res = real_llm_parse(file_path, raw_content)
        if res[0] is None:
            return
        entity_name, yaml_frontmatter, body = res

    entity_filename = f"{entity_name}.md"
    entity_filepath = os.path.join(ENTITIES_DIR, entity_filename)
    
    with open(entity_filepath, 'w', encoding='utf-8') as f:
        f.write(yaml_frontmatter + "\n" + body)
        
    print(f"[WIKI] Created semantic node at {entity_filepath}")
    update_index(entity_name, entity_filepath, format_type=".md")

def ingest_path(target_path, mock=False):
    """Router that handles either single files or massive deeply-nested directories."""
    if not os.path.exists(target_path):
        print(f"[ERR] Path not found: {target_path}")
        return

    if os.path.isfile(target_path):
        ingest_file(target_path, mock)
    elif os.path.isdir(target_path):
        print(f"[BULK INGEST] Commencing recursive deep-dive into subdirectory: {target_path}")
        for root, _, files in os.walk(target_path):
            if "archive" in root or "duplicates" in root:
                continue
            for file in files:
                if not file.startswith('.'):
                    full_path = os.path.join(root, file)
                    ingest_file(full_path, mock)

if __name__ == "__main__":
    from dotenv import load_dotenv
    load_dotenv()
    parser = argparse.ArgumentParser()
    parser.add_argument("filepath", help="The raw source file or folder to deeply ingest")
    parser.add_argument("--mock", action="store_true", help="Use mock LLM extraction")
    args = parser.parse_args()
    ingest_path(args.filepath, mock=args.mock)
